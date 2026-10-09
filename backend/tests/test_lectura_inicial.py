"""Período de lectura inicial y lectura anterior sin ceros silenciosos (cambio lectura-inicial)."""
import os
from datetime import date, datetime, timezone

from sqlalchemy import select

from app.core.security import create_access_token
from app.models.auditoria import AuditoriaLog
from app.models.boleta import BoletaMaestra
from app.models.condominio import Condominio
from app.models.lectura import LecturaParcela
from app.models.parcela import Parcela
from app.models.rol import Rol
from app.models.usuario import Usuario

MARZO = date(2099, 3, 1)


async def _cabecera(session):
    """Super admin propio de la prueba (no depende de la clave de los datos de desarrollo)."""
    rol_id = (await session.execute(select(Rol.id).where(Rol.nombre == "super_admin"))).scalar_one()
    u = Usuario(nombre="Prueba super", email=f"prueba-super-{os.urandom(3).hex()}@test.cl",
                password_hash="x", rol_id=rol_id, condominio_id=None)
    session.add(u)
    await session.flush()
    return {"Authorization": f"Bearer {create_access_token(user_id=u.id, rol='super_admin', condominio_id=None)}"}


async def _condominio(session, parcelas=("1", "2", "10")) -> Condominio:
    """Condominio nuevo, sin boletas, con parcelas activas."""
    condominio = Condominio(nombre="Condominio Lectura Inicial", rut_comunidad=f"prueba-{os.urandom(4).hex()}",
                            plan_suscripcion="basico", activo=True)
    session.add(condominio)
    await session.flush()
    session.add_all([Parcela(condominio_id=condominio.id, numero_parcela=n, activa=True) for n in parcelas])
    await session.flush()
    return condominio


async def _abrir(c, cab, condominio_id, periodo=MARZO):
    return await c.post("/api/v1/boletas/lectura-inicial", headers=cab,
                        json={"condominio_id": condominio_id, "periodo_mes": periodo.isoformat()})


async def _lecturas(c, cab, boleta_id):
    r = await c.get("/api/v1/lecturas/", headers=cab, params={"boleta_id": boleta_id})
    assert r.status_code == 200, r.text
    return r.json()


async def _tomar_todas(c, cab, boleta_id, valores: dict[int, float]):
    """Registra la lectura de cada parcela (como la captura del lector) usando el valor por parcela_id."""
    cuando = datetime(2099, 3, 2, 10, tzinfo=timezone.utc).isoformat()
    for l in await _lecturas(c, cab, boleta_id):
        r = await c.patch(f"/api/v1/lecturas/{l['id']}", headers=cab,
                          json={"lectura_actual": valores[l["parcela_id"]], "fecha_toma": cuando})
        assert r.status_code == 200, r.text


async def _crear_boleta(c, cab, condominio_id, **extra):
    return await c.post("/api/v1/boletas/", headers=cab, json={"condominio_id": condominio_id, **extra})


async def test_abre_la_lectura_inicial(tx):
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)

    r = await _abrir(c, cab, condominio.id)
    assert r.status_code == 201, r.text
    boleta = r.json()
    assert boleta["tipo"] == "lectura_inicial"
    assert boleta["items_detalle"] == []
    lecturas = await _lecturas(c, cab, boleta["id"])
    assert len(lecturas) == 3
    assert all(l["lectura_anterior"] == 0 and l["fecha_toma"] is None for l in lecturas)
    assert not any(l["lectura_anterior_editable"] for l in lecturas)
    auditoria = (await session.execute(
        select(AuditoriaLog).where(AuditoriaLog.accion == "CREATE_LECTURA_INICIAL",
                                   AuditoriaLog.condominio_id == condominio.id))).scalars().all()
    assert len(auditoria) == 1

    # Solo antes de la primera boleta
    r = await _abrir(c, cab, condominio.id, date(2099, 4, 1))
    assert r.status_code == 409


async def test_candados_de_la_lectura_inicial(tx):
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    bid = (await _abrir(c, cab, condominio.id)).json()["id"]

    intentos = [
        c.post(f"/api/v1/liquidaciones/calcular/{bid}", headers=cab),
        c.post(f"/api/v1/boletas/{bid}/validar-items", headers=cab),
        c.put(f"/api/v1/boletas/{bid}/detalles", headers=cab, json={"items_detalle": []}),
        c.post(f"/api/v1/boletas/{bid}/cerrar-liquidaciones", headers=cab),
        c.patch(f"/api/v1/boletas/{bid}", headers=cab, json={"boleta_visible_usuarios": True}),
        c.post(f"/api/v1/boletas/{bid}/imagen", headers=cab,
               files={"file": ("b.jpg", b"\xff\xd8\xff" + b"\x00" * 10, "image/jpeg")}),
        c.post(f"/api/v1/boletas/{bid}/procesar-ocr", headers=cab),
    ]
    for intento in intentos:
        r = await intento
        assert r.status_code == 409, (r.request.url, r.text)
        assert r.json()["detail"] == "El período de lectura inicial no se liquida"
    boleta = await session.get(BoletaMaestra, bid)
    await session.refresh(boleta)
    assert boleta.boleta_visible_usuarios is False

    # En la lectura inicial no hay lectura anterior que editar
    lectura = (await _lecturas(c, cab, bid))[0]
    r = await c.patch(f"/api/v1/lecturas/{lectura['id']}", headers=cab, json={"lectura_anterior": 5})
    assert r.status_code == 409


async def test_primera_boleta_hereda_la_lectura_inicial(tx):
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    inicial = (await _abrir(c, cab, condominio.id)).json()
    parcelas = {p.numero_parcela: p.id for p in (await session.execute(
        select(Parcela).where(Parcela.condominio_id == condominio.id))).scalars()}
    valores = {parcelas["1"]: 15230.0, parcelas["2"]: 8800.5, parcelas["10"]: 120.0}

    # Con la lectura inicial abierta no se puede cargar la primera boleta
    r = await _crear_boleta(c, cab, condominio.id)
    assert r.status_code == 409
    assert "lectura inicial" in r.json()["detail"]

    await _tomar_todas(c, cab, inicial["id"], valores)
    r = await c.post(f"/api/v1/boletas/{inicial['id']}/cerrar-lecturas", headers=cab)
    assert r.status_code == 200, r.text

    r = await _crear_boleta(c, cab, condominio.id)
    assert r.status_code == 201, r.text
    primera = r.json()
    assert primera["tipo"] == "regular"
    assert primera["periodo_mes"] == "2099-04-01"
    assert primera["items_detalle"] == []
    lecturas = await _lecturas(c, cab, primera["id"])
    assert {l["parcela_id"]: l["lectura_anterior"] for l in lecturas} == valores
    assert not any(l["lectura_anterior_editable"] for l in lecturas)

    # Su lectura anterior viene del período anterior: no se edita
    r = await c.patch(f"/api/v1/lecturas/{lecturas[0]['id']}", headers=cab, json={"lectura_anterior": 1})
    assert r.status_code == 409

    # La lectura inicial ya es la base de la primera boleta: no se reabre
    r = await c.post(f"/api/v1/boletas/{inicial['id']}/reabrir-lecturas", headers=cab)
    assert r.status_code == 409
    assert "período posterior" in r.json()["detail"]


async def test_reabrir_la_lectura_inicial_sin_periodo_posterior(tx):
    c, session = tx
    condominio = await _condominio(session, parcelas=("1",))
    cab = await _cabecera(session)
    inicial = (await _abrir(c, cab, condominio.id)).json()
    parcela_id = (await _lecturas(c, cab, inicial["id"]))[0]["parcela_id"]
    await _tomar_todas(c, cab, inicial["id"], {parcela_id: 100.0})
    await c.post(f"/api/v1/boletas/{inicial['id']}/cerrar-lecturas", headers=cab)
    r = await c.post(f"/api/v1/boletas/{inicial['id']}/reabrir-lecturas", headers=cab)
    assert r.status_code == 200, r.text


async def test_sin_ceros_silenciosos(tx):
    c, session = tx
    condominio = await _condominio(session)
    session.add(Parcela(condominio_id=condominio.id, numero_parcela="99", activa=False))   # inactiva: no cuenta
    await session.flush()
    cab = await _cabecera(session)

    r = await _crear_boleta(c, cab, condominio.id)
    assert r.status_code == 409
    detalle = r.json()["detail"]
    assert detalle["codigo"] == "sin_lectura_anterior"
    assert [p["numero_parcela"] for p in detalle["parcelas"]] == ["1", "2", "10"]   # orden natural
    assert (await session.execute(
        select(BoletaMaestra).where(BoletaMaestra.condominio_id == condominio.id))).first() is None

    r = await _crear_boleta(c, cab, condominio.id, aceptar_sin_lectura_anterior=True)
    assert r.status_code == 201, r.text
    lecturas = await _lecturas(c, cab, r.json()["id"])
    assert all(l["lectura_anterior"] == 0 for l in lecturas)
    assert all(l["lectura_anterior_editable"] for l in lecturas)


async def test_parcela_agregada_despues_ingresa_su_lectura_anterior(tx):
    c, session = tx
    condominio = await _condominio(session, parcelas=("1",))
    cab = await _cabecera(session)
    inicial = (await _abrir(c, cab, condominio.id)).json()
    p1 = (await _lecturas(c, cab, inicial["id"]))[0]["parcela_id"]
    await _tomar_todas(c, cab, inicial["id"], {p1: 500.0})
    await c.post(f"/api/v1/boletas/{inicial['id']}/cerrar-lecturas", headers=cab)

    nueva = Parcela(condominio_id=condominio.id, numero_parcela="2", activa=True)
    session.add(nueva)
    await session.flush()

    r = await _crear_boleta(c, cab, condominio.id)
    assert r.status_code == 409
    assert [p["numero_parcela"] for p in r.json()["detail"]["parcelas"]] == ["2"]

    r = await _crear_boleta(c, cab, condominio.id, aceptar_sin_lectura_anterior=True)
    assert r.status_code == 201, r.text
    lecturas = {l["parcela_id"]: l for l in await _lecturas(c, cab, r.json()["id"])}
    assert lecturas[p1]["lectura_anterior_editable"] is False
    assert lecturas[nueva.id]["lectura_anterior_editable"] is True

    r = await c.patch(f"/api/v1/lecturas/{lecturas[nueva.id]['id']}", headers=cab,
                      json={"lectura_anterior": 300.0, "lectura_actual": 420.0})
    assert r.status_code == 200, r.text
    assert r.json()["kwh_consumidos"] == 120.0


# ---- Carga masiva desde Excel ---------------------------------------------------------------------------

def _xlsx(filas, encabezados=("Parcela", "Propietario", "Lectura inicial")) -> bytes:
    import io
    from openpyxl import Workbook
    libro = Workbook()
    hoja = libro.active
    hoja.append(["Lecturas iniciales del condominio"])   # título antes de los encabezados
    hoja.append(list(encabezados))
    for fila in filas:
        hoja.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


async def _importar(c, cab, boleta_id, contenido, aplicar=False):
    return await c.post(f"/api/v1/boletas/{boleta_id}/lecturas-iniciales/importar", headers=cab,
                        params={"aplicar": str(aplicar).lower()},
                        files={"archivo": ("lecturas.xlsx", contenido,
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


async def test_plantilla_se_lee_de_vuelta(tx):
    import io
    from openpyxl import load_workbook
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    bid = (await _abrir(c, cab, condominio.id)).json()["id"]
    r = await c.get(f"/api/v1/boletas/{bid}/lecturas-iniciales/plantilla", headers=cab)
    assert r.status_code == 200, r.text
    filas = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert filas[0] == ("Parcela", "Propietario", "Lectura inicial", "Saldo luz")
    assert [f[0] for f in filas[1:]] == ["1", "2", "10"]   # orden natural

    # La misma plantilla, completada, sirve para importar
    completada = _xlsx([(f[0], f[1], 100 * (n + 1)) for n, f in enumerate(filas[1:])])
    r = await _importar(c, cab, bid, completada)
    assert r.status_code == 200, r.text
    assert len(r.json()["a_aplicar"]) == 3


async def test_vista_previa_no_modifica_y_aplicar_si(tx):
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    bid = (await _abrir(c, cab, condominio.id)).json()["id"]
    planilla = _xlsx([("Parcela 1", "", 15230), (2.0, "", "8.800,5"), ("10", "", "120.500"), ("", "", None)])

    r = await _importar(c, cab, bid, planilla)
    assert r.status_code == 200, r.text
    vista = r.json()
    assert vista["errores"] == [] and vista["aplicadas"] == 0
    assert sorted((f["numero_parcela"], f["valor"]) for f in vista["a_aplicar"]) == [
        ("1", 15230.0), ("10", 120500.0), ("2", 8800.5)]
    assert all(l["fecha_toma"] is None for l in await _lecturas(c, cab, bid))   # nada aplicado

    r = await _importar(c, cab, bid, planilla, aplicar=True)
    assert r.status_code == 200, r.text
    assert r.json()["aplicadas"] == 3
    lecturas = await _lecturas(c, cab, bid)
    assert all(l["fecha_toma"] for l in lecturas)
    auditoria = (await session.execute(select(AuditoriaLog).where(
        AuditoriaLog.accion == "IMPORTAR_LECTURAS_INICIALES", AuditoriaLog.condominio_id == condominio.id))).scalars().all()
    assert len(auditoria) == 1

    # Reaplicar lo mismo: nada cambia; un valor distinto reemplaza la lectura tomada
    r = await _importar(c, cab, bid, _xlsx([("1", "", 15230), ("2", "", 9000)]))
    vista = r.json()
    assert vista["sin_cambio"] == 1
    assert [(f["numero_parcela"], f["reemplaza"], f["valor_actual"]) for f in vista["a_aplicar"]] == [("2", True, 8800.5)]


async def test_errores_no_aplican_nada(tx):
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    bid = (await _abrir(c, cab, condominio.id)).json()["id"]
    planilla = _xlsx([("1", "", 100), ("99", "", 5), ("1", "", 7), ("2", "", -3), ("10", "", "abc"), ("2", "", None)],
                     encabezados=("Unidad", "Nombre", "Lectura"))
    r = await _importar(c, cab, bid, planilla)
    errores = r.json()["errores"]
    # 99 no existe, 1 repetida, -3 negativa, "abc" no es número, 2 repetida (aunque venga vacía)
    assert [e["fila"] for e in errores] == [4, 5, 6, 7, 8]
    r = await _importar(c, cab, bid, planilla, aplicar=True)
    assert r.status_code == 422
    assert all(l["fecha_toma"] is None for l in await _lecturas(c, cab, bid))


async def test_importar_solo_en_lectura_inicial_abierta(tx):
    c, session = tx
    condominio = await _condominio(session, parcelas=("1",))
    cab = await _cabecera(session)
    inicial = (await _abrir(c, cab, condominio.id)).json()
    planilla = _xlsx([("1", "", 100)])

    # Archivo que no es planilla, o sin las columnas
    assert (await _importar(c, cab, inicial["id"], b"no es un excel")).status_code == 422
    assert (await _importar(c, cab, inicial["id"], _xlsx([("1", 2)], encabezados=("A", "B")))).status_code == 422

    await _importar(c, cab, inicial["id"], planilla, aplicar=True)
    await c.post(f"/api/v1/boletas/{inicial['id']}/cerrar-lecturas", headers=cab)
    assert (await _importar(c, cab, inicial["id"], planilla)).status_code == 409   # cerrada

    regular = (await _crear_boleta(c, cab, condominio.id)).json()
    assert (await _importar(c, cab, regular["id"], planilla)).status_code == 409   # período regular
    r = await c.get(f"/api/v1/boletas/{regular['id']}/lecturas-iniciales/plantilla", headers=cab)
    assert r.status_code == 409


async def test_minimo_para_liquidar_sin_items_ni_imagen(tx):
    """Primera boleta sin boleta física: solo los tres totales, sin ítems ni imagen, ya permite liquidar."""
    c, session = tx
    condominio = await _condominio(session)
    cab = await _cabecera(session)
    inicial = (await _abrir(c, cab, condominio.id)).json()
    parcelas = {l["parcela_id"] for l in await _lecturas(c, cab, inicial["id"])}
    await _tomar_todas(c, cab, inicial["id"], {p: 1000.0 for p in parcelas})
    await c.post(f"/api/v1/boletas/{inicial['id']}/cerrar-lecturas", headers=cab)

    octubre = (await _crear_boleta(c, cab, condominio.id)).json()
    consumos = dict(zip(sorted(parcelas), (100.0, 200.0, 300.0)))
    await _tomar_todas(c, cab, octubre["id"], {p: 1000.0 + kwh for p, kwh in consumos.items()})

    # Sin total emisión no se puede calcular (antes era un error 500)
    r = await c.put(f"/api/v1/boletas/{octubre['id']}/detalles", headers=cab,
                    json={"total_kwh_compania": 700, "monto_neto_electricidad_consumida": 100000, "items_detalle": []})
    assert r.status_code == 200, r.text
    assert (await c.post(f"/api/v1/boletas/{octubre['id']}/validar-items", headers=cab)).status_code == 200
    r = await c.post(f"/api/v1/liquidaciones/calcular/{octubre['id']}", headers=cab)
    assert r.status_code == 422
    assert "monto_total_emision" in r.json()["detail"]

    # Con los tres totales y ningún ítem: se liquida y el total cuadra exacto con la emisión
    await c.put(f"/api/v1/boletas/{octubre['id']}/detalles", headers=cab,
                json={"total_kwh_compania": 700, "monto_neto_electricidad_consumida": 100000,
                      "monto_total_emision": 119000, "items_detalle": []})
    await c.post(f"/api/v1/boletas/{octubre['id']}/validar-items", headers=cab)
    r = await c.post(f"/api/v1/liquidaciones/calcular/{octubre['id']}", headers=cab)
    assert r.status_code == 200, r.text
    liquidaciones = r.json()
    assert len(liquidaciones) == 3
    assert sum(l["total_pagar_mes"] for l in liquidaciones) == 119000   # cuadre al peso
