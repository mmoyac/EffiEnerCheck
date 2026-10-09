"""
Cuenta corriente de luz y cobranza (cambio cobranza-energia): saldo inicial, abonos de cualquier monto
imputados a la deuda más antigua, anulación auditada, resumen y vista del comunero.
"""
import io
import os
from datetime import date

from openpyxl import Workbook
from sqlalchemy import insert, select

from app.core.security import create_access_token
from app.models.auditoria import AuditoriaLog
from app.models.boleta import BoletaMaestra
from app.models.condominio import Condominio
from app.models.condominio_modulo import CondominioModulo
from app.models.cuenta_luz import MovimientoLuz
from app.models.lectura import LecturaParcela
from app.models.liquidacion import LiquidacionParcela
from app.models.parcela import Parcela
from app.models.rol import Rol
from app.models.usuario import Usuario
from app.models.usuario_parcela import usuario_parcelas

HOY = date.today()


async def _usuario(session, rol, condominio_id):
    rol_id = (await session.execute(select(Rol.id).where(Rol.nombre == rol))).scalar_one()
    u = Usuario(nombre=f"Prueba {rol}", email=f"cuenta-{rol}-{os.urandom(3).hex()}@test.cl", password_hash="x",
                rol_id=rol_id, condominio_id=condominio_id)
    session.add(u)
    await session.flush()
    return u, {"Authorization": f"Bearer {create_access_token(user_id=u.id, rol=rol, condominio_id=condominio_id)}"}


async def _escenario(session, saldo_inicial=None):
    """Parcelas 2 y 3. Octubre y noviembre publicados; diciembre cerrado sin publicar."""
    c = Condominio(nombre="Condominio Cuenta", rut_comunidad=f"cuenta-{os.urandom(4).hex()}",
                   plan_suscripcion="basico", activo=True,
                   modulos=[CondominioModulo(modulo="portal"), CondominioModulo(modulo="energia")])
    session.add(c)
    await session.flush()
    admin, cab = await _usuario(session, "admin_condominio", c.id)
    p2 = Parcela(condominio_id=c.id, numero_parcela="2", propietario_nombre="Ana", activa=True)
    p3 = Parcela(condominio_id=c.id, numero_parcela="3", propietario_nombre="Beto", activa=True)
    session.add_all([p2, p3])
    await session.flush()
    liqs = {}
    for mes, publicado, (m2, m3) in ((10, True, (10_000, 30_000)), (11, True, (11_000, 31_000)), (12, False, (12_000, 32_000))):
        b = BoletaMaestra(condominio_id=c.id, periodo_mes=date(2026, mes, 1), creado_por=admin.id,
                          estado="publicada" if publicado else "validada", lecturas_cerradas=True,
                          liquidaciones_cerradas=True, boleta_visible_usuarios=publicado)
        session.add(b)
        await session.flush()
        for p, monto in ((p2, m2), (p3, m3)):
            liqs[(mes, p.numero_parcela)] = LiquidacionParcela(
                parcela_id=p.id, boleta_id=b.id, monto_energia_kwh=monto, monto_prorrateo_variable=0,
                monto_cuota_fija=0, total_pagar_mes=monto)
            session.add(liqs[(mes, p.numero_parcela)])
    if saldo_inicial:
        session.add(MovimientoLuz(condominio_id=c.id, parcela_id=p2.id, tipo="saldo_inicial", monto=saldo_inicial,
                                  fecha=date(2026, 9, 1), creado_por=admin.id))
    await session.flush()
    return c, admin, cab, p2, p3, liqs


async def _abonar(c, cab, items, fecha=HOY, nota="Gasto común"):
    return await c.post("/api/v1/cuenta-luz/abonos", headers=cab, json={
        "fecha": fecha.isoformat(), "nota": nota, "items": [{"parcela_id": p, "monto": m} for p, m in items]})


async def test_abono_parcial_se_imputa_a_lo_mas_antiguo(tx):
    c, session = tx
    _, _, cab, p2, _, liqs = await _escenario(session, saldo_inicial=5_000)
    # Deuda de la parcela 2: saldo inicial 5.000 + oct 10.000 + nov 11.000 = 26.000 (dic no está publicado)
    r = await c.get(f"/api/v1/cuenta-luz/parcelas/{p2.id}", headers=cab)
    assert r.json()["saldo"] == 26_000

    r = await _abonar(c, cab, [(p2.id, 12_000)])
    assert r.status_code == 201, r.text
    cuenta = r.json()[0]
    assert cuenta["saldo"] == 14_000
    # 12.000 cubren el saldo inicial (5.000) y 7.000 de octubre
    assert [(x["tipo"], x["estado"], x["cubierto"]) for x in cuenta["cargos"]] == [
        ("saldo_inicial", "pagado", 5_000), ("mes", "parcial", 7_000), ("mes", "pendiente", 0)]
    await session.refresh(liqs[(10, "2")])
    assert liqs[(10, "2")].monto_abonado == 7_000 and liqs[(10, "2")].pagado is False

    # Otro abono completa octubre y deja noviembre parcial
    cuenta = (await _abonar(c, cab, [(p2.id, 5_000)])).json()[0]
    assert [x["estado"] for x in cuenta["cargos"]] == ["pagado", "pagado", "parcial"]
    await session.refresh(liqs[(10, "2")])
    assert liqs[(10, "2")].pagado is True and liqs[(10, "2")].fecha_pago is not None


async def test_abono_queda_auditado_y_se_anula_sin_borrarse(tx):
    c, session = tx
    condominio, _, cab, p2, _, _ = await _escenario(session)
    await _abonar(c, cab, [(p2.id, 8_000)])
    reg = (await session.execute(select(AuditoriaLog).where(
        AuditoriaLog.accion == "REGISTRAR_ABONO_LUZ", AuditoriaLog.condominio_id == condominio.id))).scalar_one()
    assert reg.detalles["saldo_antes"] == 21_000 and reg.detalles["saldo_despues"] == 13_000

    abono = (await session.execute(select(MovimientoLuz).where(MovimientoLuz.parcela_id == p2.id))).scalar_one()
    r = await c.post(f"/api/v1/cuenta-luz/abonos/{abono.id}/anular", headers=cab, json={"motivo": "Registrado en la parcela equivocada"})
    assert r.status_code == 200, r.text
    cuenta = r.json()
    assert cuenta["saldo"] == 21_000
    assert cuenta["abonos"][0]["anulado"] is True and cuenta["abonos"][0]["motivo_anulacion"]   # sigue visible
    assert (await session.get(MovimientoLuz, abono.id)) is not None
    anul = (await session.execute(select(AuditoriaLog).where(
        AuditoriaLog.accion == "ANULAR_ABONO_LUZ", AuditoriaLog.condominio_id == condominio.id))).scalar_one()
    assert anul.detalles["saldo_despues"] == 21_000
    # No se anula dos veces, y el motivo es obligatorio
    assert (await c.post(f"/api/v1/cuenta-luz/abonos/{abono.id}/anular", headers=cab,
                         json={"motivo": "otra vez"})).status_code == 409
    assert (await c.post(f"/api/v1/cuenta-luz/abonos/{abono.id}/anular", headers=cab,
                         json={"motivo": ""})).status_code == 422


async def test_resumen_de_cobranza(tx):
    c, session = tx
    _, _, cab, p2, p3, _ = await _escenario(session, saldo_inicial=5_000)
    await _abonar(c, cab, [(p2.id, 26_000), (p3.id, 40_000)])   # p2 al día; p3 debe 21.000
    r = await c.get("/api/v1/cuenta-luz/resumen", headers=cab)
    assert r.status_code == 200, r.text
    datos = r.json()
    assert datos["totales"] == {"saldo_inicial": 5_000, "emitido": 82_000, "cargos": 87_000,
                                "abonado": 66_000, "por_cobrar": 21_000, "a_favor": 0}
    assert [d["numero_parcela"] for d in datos["deudores"]] == ["3"]
    deudor = datos["deudores"][0]
    assert deudor["saldo"] == 21_000 and [x["estado"] for x in deudor["pendientes"]] == ["parcial"]
    # Por período: noviembre primero, octubre y el saldo inicial al final (diciembre no publicado no aparece)
    assert [(p["tipo"], p["fecha"]) for p in datos["periodos"]] == [
        ("mes", "2026-11-01"), ("mes", "2026-10-01"), ("saldo_inicial", "2026-09-01")]


async def test_saldo_a_favor_se_aplica_al_publicar(tx):
    c, session = tx
    _, admin, cab, p2, _, liqs = await _escenario(session)
    await _abonar(c, cab, [(p2.id, 30_000)])      # deuda 21.000 → 9.000 a favor
    assert (await c.get(f"/api/v1/cuenta-luz/parcelas/{p2.id}", headers=cab)).json()["saldo"] == -9_000
    diciembre = await session.get(BoletaMaestra, liqs[(12, "2")].boleta_id)
    r = await c.patch(f"/api/v1/boletas/{diciembre.id}", headers=cab, json={"boleta_visible_usuarios": True})
    assert r.status_code == 200, r.text
    await session.refresh(liqs[(12, "2")])
    assert liqs[(12, "2")].monto_abonado == 9_000 and liqs[(12, "2")].pagado is False   # 9.000 de 12.000


async def test_validaciones_de_abono(tx):
    c, session = tx
    _, _, cab, p2, _, _ = await _escenario(session)
    _, _, _, ajena, _, _ = await _escenario(session)
    assert (await _abonar(c, cab, [(ajena.id, 1_000)])).status_code == 404
    assert (await _abonar(c, cab, [(p2.id, 0)])).status_code == 422
    assert (await _abonar(c, cab, [(p2.id, 100), (p2.id, 200)])).status_code == 422
    assert (await _abonar(c, cab, [(p2.id, 100)], fecha=date(2999, 1, 1))).status_code == 422


async def test_comunero_ve_solo_su_cuenta(tx):
    c, session = tx
    condominio, _, cab, p2, p3, _ = await _escenario(session, saldo_inicial=5_000)
    comunero, cab_comunero = await _usuario(session, "comunero", condominio.id)
    await session.execute(insert(usuario_parcelas).values(usuario_id=comunero.id, parcela_id=p2.id))
    await session.flush()
    r = await c.get(f"/api/v1/cuenta-luz/parcelas/{p2.id}", headers=cab_comunero)
    assert r.status_code == 200 and r.json()["saldo"] == 26_000
    assert (await c.get(f"/api/v1/cuenta-luz/parcelas/{p3.id}", headers=cab_comunero)).status_code == 404
    assert (await c.get("/api/v1/cuenta-luz/resumen", headers=cab_comunero)).status_code == 403
    assert (await _abonar(c, cab_comunero, [(p2.id, 1_000)])).status_code == 403


def _planilla(filas, encabezados=("Parcela", "Propietario", "Lectura inicial", "Saldo luz")) -> bytes:
    libro = Workbook()
    libro.active.append(list(encabezados))
    for f in filas:
        libro.active.append(list(f))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


async def test_saldo_inicial_desde_la_planilla_de_onboarding(tx):
    c, session = tx
    cond = Condominio(nombre="Condominio Onboarding", rut_comunidad=f"onb-{os.urandom(4).hex()}",
                      plan_suscripcion="basico", activo=True,
                      modulos=[CondominioModulo(modulo="portal"), CondominioModulo(modulo="energia")])
    session.add(cond)
    await session.flush()
    _, cab = await _usuario(session, "admin_condominio", cond.id)
    p1 = Parcela(condominio_id=cond.id, numero_parcela="1", activa=True)
    p2 = Parcela(condominio_id=cond.id, numero_parcela="2", activa=True)
    session.add_all([p1, p2])
    await session.flush()
    bid = (await c.post("/api/v1/boletas/lectura-inicial", headers=cab, json={"periodo_mes": "2026-09-01"})).json()["id"]
    url = f"/api/v1/boletas/{bid}/lecturas-iniciales/importar"
    archivo = lambda filas: {"archivo": ("x.xlsx", _planilla(filas), "application/octet-stream")}

    r = await c.post(url, headers=cab, files=archivo([("1", "", 100, "45.000"), ("2", "", 200, None)]))
    vista = r.json()
    assert [(s["numero_parcela"], s["valor"], s["valor_actual"]) for s in vista["saldos"]] == [("1", 45_000, None)]
    assert (await c.post(url, headers=cab, params={"aplicar": "true"},
                         files=archivo([("1", "", 100, "45.000"), ("2", "", 200, None)]))).status_code == 200
    assert (await c.get(f"/api/v1/cuenta-luz/parcelas/{p1.id}", headers=cab)).json()["saldo"] == 45_000

    # Corregir el saldo no lo pisa: el anterior queda anulado con motivo
    await c.post(url, headers=cab, params={"aplicar": "true"}, files=archivo([("1", "", 100, 40000)]))
    movs = (await session.execute(select(MovimientoLuz).where(MovimientoLuz.parcela_id == p1.id)
                                  .order_by(MovimientoLuz.id))).scalars().all()
    assert [(m.monto, m.anulado) for m in movs] == [(45_000, True), (40_000, False)]
    assert movs[0].motivo_anulacion

    # Saldo negativo: error de fila, nada se aplica
    r = await c.post(url, headers=cab, params={"aplicar": "true"}, files=archivo([("2", "", 200, -5)]))
    assert r.status_code == 422

    # La plantilla trae el saldo vigente
    plantilla = await c.get(f"/api/v1/boletas/{bid}/lecturas-iniciales/plantilla", headers=cab)
    from openpyxl import load_workbook
    filas = list(load_workbook(io.BytesIO(plantilla.content)).active.iter_rows(values_only=True))
    assert filas[1][3] == 40_000 and filas[2][3] is None
