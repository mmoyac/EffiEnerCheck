"""
Carga masiva de lecturas iniciales desde Excel (cambio lectura-inicial).

- `plantilla()`: genera el .xlsx con una fila por parcela del período (Parcela, Propietario, Lectura inicial).
- `leer_planilla()`: interpreta el .xlsx subido y lo cruza con las lecturas del período. No toca la base:
  devuelve lo que se aplicaría, lo que no cambia, las filas vacías y los errores por fila.
"""
import io
import re
import unicodedata
import warnings
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

MAX_BYTES = 2 * 1024 * 1024
ENCABEZADOS_PARCELA = ("parcela", "unidad")
ENCABEZADOS_LECTURA = ("lectura inicial", "lectura")


class PlanillaInvalida(ValueError):
    """El archivo no es una planilla legible o no tiene las columnas esperadas."""


@dataclass
class FilaAplicable:
    lectura_id: int
    parcela_id: int
    numero_parcela: str
    valor: float
    valor_actual: float | None   # None si aún no se tomaba
    reemplaza: bool              # ya había una lectura tomada con otro valor


@dataclass
class Resultado:
    a_aplicar: list[FilaAplicable] = field(default_factory=list)
    sin_cambio: int = 0
    vacias: int = 0
    errores: list[dict] = field(default_factory=list)


def _normalizar(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor if valor is not None else "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def clave_parcela(valor) -> str:
    """'Parcela 23', ' 23 ', 23.0 y '23' son la misma parcela."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = _normalizar(valor)
    texto = re.sub(r"^parcela\s*", "", texto)
    return texto.replace(" ", "")


def interpretar_valor(valor) -> float:
    """Número de la celda. Texto en formato chileno: '15.230' = 15230 (miles), '15.230,5' = 15230.5."""
    if isinstance(valor, bool):
        raise ValueError("no es un número")
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
        texto = texto.replace(".", "")
    try:
        return float(texto)
    except ValueError:
        raise ValueError("no es un número") from None


def plantilla(filas: list[tuple[str, str | None, float | None]], titulo: str) -> bytes:
    """filas: (numero_parcela, propietario, lectura_inicial_o_None), ya en orden natural."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Lecturas iniciales"
    hoja.append(["Parcela", "Propietario", "Lectura inicial"])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for numero, propietario, lectura in filas:
        hoja.append([numero, propietario or "", lectura])
    hoja.column_dimensions["A"].width = 12
    hoja.column_dimensions["B"].width = 36
    hoja.column_dimensions["C"].width = 16
    hoja.freeze_panes = "A2"
    libro.properties.title = titulo
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def leer_planilla(contenido: bytes, lecturas: list[tuple[int, int, str, float, bool]]) -> Resultado:
    """
    lecturas del período: (lectura_id, parcela_id, numero_parcela, lectura_actual, tomada).
    Lanza PlanillaInvalida si el archivo no se puede leer o le faltan columnas.
    """
    if len(contenido) > MAX_BYTES:
        raise PlanillaInvalida("La planilla supera los 2 MB")
    if not contenido.startswith(b"PK"):
        raise PlanillaInvalida("El archivo no es una planilla Excel (.xlsx)")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
            filas = list(hoja.iter_rows(values_only=True))
    except Exception:
        raise PlanillaInvalida("No se pudo leer la planilla. Guárdala como .xlsx e intenta de nuevo.") from None

    col_parcela = col_lectura = None
    inicio = 0
    for n, fila in enumerate(filas):
        encabezados = [_normalizar(c) for c in fila]
        col_parcela = next((encabezados.index(h) for h in ENCABEZADOS_PARCELA if h in encabezados), None)
        col_lectura = next((encabezados.index(h) for h in ENCABEZADOS_LECTURA if h in encabezados), None)
        if col_parcela is not None and col_lectura is not None:
            inicio = n + 1
            break
    if col_parcela is None or col_lectura is None:
        raise PlanillaInvalida("La planilla debe tener las columnas «Parcela» y «Lectura inicial»")

    por_clave = {clave_parcela(numero): (lid, pid, numero, actual, tomada)
                 for lid, pid, numero, actual, tomada in lecturas}
    resultado = Resultado()
    vistas: dict[str, int] = {}
    for n, fila in enumerate(filas[inicio:], start=inicio + 1):
        celda_parcela = fila[col_parcela] if col_parcela < len(fila) else None
        celda_lectura = fila[col_lectura] if col_lectura < len(fila) else None
        if celda_parcela in (None, "") and celda_lectura in (None, ""):
            continue
        clave = clave_parcela(celda_parcela)
        if not clave:
            resultado.errores.append({"fila": n, "mensaje": "Falta el número de parcela"})
            continue
        if clave not in por_clave:
            resultado.errores.append({"fila": n, "mensaje": f"La parcela «{celda_parcela}» no está en el período"})
            continue
        if clave in vistas:
            resultado.errores.append({"fila": n, "mensaje": f"La parcela «{celda_parcela}» se repite (ya está en la fila {vistas[clave]})"})
            continue
        vistas[clave] = n
        if celda_lectura is None or str(celda_lectura).strip() == "":
            resultado.vacias += 1
            continue
        try:
            valor = interpretar_valor(celda_lectura)
        except ValueError:
            resultado.errores.append({"fila": n, "mensaje": f"La lectura «{celda_lectura}» no es un número"})
            continue
        if valor < 0:
            resultado.errores.append({"fila": n, "mensaje": f"La lectura no puede ser negativa ({valor:g})"})
            continue
        lid, pid, numero, actual, tomada = por_clave[clave]
        if tomada and abs(actual - valor) < 1e-9:
            resultado.sin_cambio += 1
            continue
        resultado.a_aplicar.append(FilaAplicable(
            lectura_id=lid, parcela_id=pid, numero_parcela=numero, valor=valor,
            valor_actual=actual if tomada else None, reemplaza=tomada,
        ))
    return resultado
