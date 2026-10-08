"""
Carga el padrón de un condominio desde la planilla de residentes (MATRIZ RESIDENTES.xlsx).

    python -m app.db.cargar_residentes <planilla.xlsx> --condominio "Santa Laura" --rut <rut> [--simular]

Crea, si no existen:
  - el condominio (por nombre; --rut solo se usa al crearlo),
  - una parcela por cada "Unidad" (propietario = primer "Dueño" de la unidad),
  - un usuario comunero por cada "Correo", con su "Teléfono" normalizado, asignado a su unidad.

Los residentes nuevos quedan como cuentas pendientes, sin clave: cada uno crea la suya con su
invitación (spec acceso-por-enlace). Es idempotente: no toca las cuentas que ya existen, solo
completa el teléfono si estaba vacío y agrega las asignaciones que falten.
Con --simular muestra el resumen sin escribir nada.

La planilla tiene datos personales: nunca va al repositorio ni a la imagen (ver DEPLOY.md).
"""
import argparse
import asyncio
import sys
import unicodedata
import warnings
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import insert, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.audit import registrar_auditoria
from app.db.session import AsyncSessionLocal, engine
from app.models.condominio import Condominio
from app.models.parcela import Parcela
from app.models.rol import Rol
from app.models.usuario import Usuario
from app.models.usuario_parcela import usuario_parcelas
from app.utils.telefono import normalizar_telefono

import app.db.base  # noqa: F401 — registra todos los modelos (relaciones por nombre)

# Encabezado normalizado → campo. La fila de encabezados se busca por la celda "Unidad".
COLUMNAS = {
    "unidad": "unidad",
    "nombre y apellido": "nombre",
    "rol": "rol",
    "correo": "email",
    "telefono": "telefono",
}


@dataclass
class Residente:
    fila: int
    unidad: str
    nombre: str
    rol: str
    email: str
    telefono: str | None


@dataclass
class Resumen:
    condominio_creado: bool = False
    parcelas_creadas: int = 0
    usuarios_creados: int = 0
    telefonos_completados: int = 0
    asignaciones_creadas: int = 0
    avisos: list[str] = field(default_factory=list)


def _normalizar_encabezado(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def _texto(valor) -> str:
    return " ".join(str(valor).split()) if valor is not None else ""


def leer_planilla(ruta: Path) -> tuple[list[Residente], list[str]]:
    with warnings.catch_warnings():
        # Las planillas exportadas traen extensiones de Excel que openpyxl ignora sin problema
        warnings.simplefilter("ignore", UserWarning)
        hoja = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))

    indices: dict[str, int] | None = None
    inicio = 0
    for n, fila in enumerate(filas):
        encabezados = [_normalizar_encabezado(c) for c in fila]
        if "unidad" in encabezados:
            indices = {campo: encabezados.index(h) for h, campo in COLUMNAS.items() if h in encabezados}
            inicio = n + 1
            break
    faltan = set(COLUMNAS.values()) - set(indices or {})
    if indices is None or faltan:
        raise ValueError(f"La planilla no tiene las columnas esperadas: faltan {sorted(faltan)}")

    residentes: list[Residente] = []
    avisos: list[str] = []
    for n, fila in enumerate(filas[inicio:], start=inicio + 1):
        unidad = _texto(fila[indices["unidad"]])
        email = _texto(fila[indices["email"]]).lower()
        if not unidad and not email:
            continue
        if not unidad or not email:
            avisos.append(f"fila {n}: sin unidad o sin correo, se omite")
            continue
        try:
            telefono = normalizar_telefono(fila[indices["telefono"]])
        except ValueError:
            avisos.append(f"fila {n} ({email}): teléfono inválido, queda vacío")
            telefono = None
        residentes.append(Residente(
            fila=n,
            unidad=unidad,
            nombre=_texto(fila[indices["nombre"]]) or email,
            rol=_texto(fila[indices["rol"]]),
            email=email,
            telefono=telefono,
        ))
    return residentes, avisos


def _orden_natural(unidad: str) -> tuple:
    cabeza = unidad.split()[0]
    return (int(cabeza) if cabeza.isdigit() else float("inf"), unidad)


async def cargar(
    db: AsyncSession,
    residentes: list[Residente],
    *,
    condominio_nombre: str,
    rut: str | None,
) -> Resumen:
    """Aplica la carga en la sesión, sin commit: quien llama confirma o descarta."""
    resumen = Resumen()

    condominio = (await db.execute(
        select(Condominio).where(Condominio.nombre == condominio_nombre)
    )).scalar_one_or_none()
    if condominio is None:
        if not rut:
            raise ValueError(f"No existe el condominio '{condominio_nombre}': indica --rut para crearlo")
        condominio = Condominio(nombre=condominio_nombre, rut_comunidad=rut, plan_suscripcion="premium", activo=True)
        db.add(condominio)
        await db.flush()
        resumen.condominio_creado = True

    # ---- Parcelas ------------------------------------------------------------------------------------
    parcelas = {
        p.numero_parcela: p
        for p in (await db.execute(select(Parcela).where(Parcela.condominio_id == condominio.id))).scalars()
    }
    # Propietario de la parcela: el primer "Dueño" de la unidad o, si no hay, el primer residente.
    propietarios: dict[str, str] = {}
    for r in sorted(residentes, key=lambda r: r.rol.lower() != "dueño"):
        propietarios.setdefault(r.unidad, r.nombre)
    for unidad in sorted({r.unidad for r in residentes}, key=_orden_natural):
        if unidad in parcelas:
            continue
        parcela = Parcela(
            condominio_id=condominio.id,
            numero_parcela=unidad,
            propietario_nombre=propietarios.get(unidad),
            activa=True,
        )
        db.add(parcela)
        parcelas[unidad] = parcela
        resumen.parcelas_creadas += 1
    await db.flush()

    # ---- Usuarios y asignaciones --------------------------------------------------------------------
    rol_comunero = (await db.execute(select(Rol).where(Rol.nombre == "comunero"))).scalar_one()
    for r in residentes:
        usuario = (await db.execute(
            select(Usuario).options(selectinload(Usuario.rol)).where(Usuario.email == r.email)
        )).scalar_one_or_none()

        if usuario is None:
            usuario = Usuario(
                nombre=r.nombre,
                email=r.email,
                password_hash=None,   # pendiente: crea su clave con la invitación
                rol_id=rol_comunero.id,
                condominio_id=condominio.id,
                telefono=r.telefono,
            )
            db.add(usuario)
            await db.flush()
            resumen.usuarios_creados += 1
        elif usuario.rol.nombre != "comunero" or usuario.condominio_id != condominio.id:
            resumen.avisos.append(
                f"fila {r.fila} ({r.email}): el correo ya es de una cuenta {usuario.rol.nombre} "
                "de otro condominio o rol; no se asigna la parcela"
            )
            continue
        elif not usuario.telefono and r.telefono:
            usuario.telefono = r.telefono
            resumen.telefonos_completados += 1

        parcela = parcelas[r.unidad]
        asignada = (await db.execute(select(usuario_parcelas).where(
            usuario_parcelas.c.usuario_id == usuario.id,
            usuario_parcelas.c.parcela_id == parcela.id,
        ))).first()
        if not asignada:
            await db.execute(insert(usuario_parcelas).values(usuario_id=usuario.id, parcela_id=parcela.id))
            resumen.asignaciones_creadas += 1

    super_admin_id = (await db.execute(
        select(Usuario.id).join(Rol).where(Rol.nombre == "super_admin").order_by(Usuario.id).limit(1)
    )).scalar_one_or_none()
    if super_admin_id is not None:
        await registrar_auditoria(
            db,
            usuario_id=super_admin_id,
            condominio_id=condominio.id,
            accion="RESIDENTES_CARGADOS",
            detalles={
                "parcelas_creadas": resumen.parcelas_creadas,
                "usuarios_creados": resumen.usuarios_creados,
                "asignaciones_creadas": resumen.asignaciones_creadas,
            },
        )
    return resumen



async def main(args: argparse.Namespace) -> None:
    residentes, avisos = leer_planilla(Path(args.planilla))
    print(f"Planilla: {len(residentes)} residentes en {len({r.unidad for r in residentes})} unidades")

    async with AsyncSessionLocal() as db:
        resumen = await cargar(db, residentes, condominio_nombre=args.condominio, rut=args.rut)
        if args.simular:
            await db.rollback()
        else:
            await db.commit()
    await engine.dispose()

    print(("SIMULACIÓN (no se guardó nada) — " if args.simular else "") +
          f"condominio {'creado' if resumen.condominio_creado else 'existente'} | "
          f"parcelas nuevas: {resumen.parcelas_creadas} | usuarios nuevos: {resumen.usuarios_creados} | "
          f"teléfonos completados: {resumen.telefonos_completados} | asignaciones nuevas: {resumen.asignaciones_creadas}")
    for aviso in avisos + resumen.avisos:
        print(f"  AVISO {aviso}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Carga parcelas y residentes desde la planilla de residentes")
    parser.add_argument("planilla", help="Ruta del .xlsx (dentro del contenedor, p. ej. /tmp/residentes.xlsx)")
    parser.add_argument("--condominio", required=True, help='Nombre del condominio, p. ej. "Santa Laura"')
    parser.add_argument("--rut", help="RUT de la comunidad (solo si el condominio no existe)")
    parser.add_argument("--simular", action="store_true", help="Muestra el resumen sin guardar")
    try:
        asyncio.run(main(parser.parse_args()))
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
